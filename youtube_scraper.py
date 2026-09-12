#!/usr/bin/env python3
"""
YouTube channel scraper → Excel (.xlsx)
Запуск: py youtube_scraper.py [--url КАНАЛ] [--limit N] [--output ФАЙЛ]
"""

import sys
import subprocess
import re
import json
import os
import glob
import time
import argparse
import platform
from datetime import datetime
from pathlib import Path

if sys.stdout.encoding and sys.stdout.encoding.lower() != "utf-8":
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")


# ─── автоустановка зависимостей ────────────────────────────────────────────
def ensure_deps():
    missing = []
    try:
        import yt_dlp  # noqa: F401
    except ImportError:
        missing.append("yt-dlp")
    try:
        import openpyxl  # noqa: F401
    except ImportError:
        missing.append("openpyxl")
    if missing:
        print(f"Устанавливаю: {', '.join(missing)}...")
        subprocess.check_call([sys.executable, "-m", "pip", "install", "-q"] + missing)


ensure_deps()
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


# ─── открытие папки кроссплатформенно ───────────────────────────────────────
def open_file_in_explorer(filepath: str):
    """Открывает папку и выделяет файл (Windows/Mac/Linux)."""
    filepath = os.path.abspath(filepath)
    
    try:
        system = platform.system()
        
        if system == "Windows":
            # Windows: используем explorer с /select
            subprocess.Popen(["explorer", "/select,", filepath])
        elif system == "Darwin":
            # macOS: используем open с -R (reveal)
            subprocess.Popen(["open", "-R", filepath])
        else:
            # Linux: открываем папку в файловом менеджере
            folder = os.path.dirname(filepath)
            subprocess.Popen(["xdg-open", folder])
        
        print(f"✓ Папка открыта: {os.path.dirname(filepath)}")
    except Exception as e:
        print(f"⚠ Не удалось открыть папку: {e}")
        print(f"  Файл находится здесь: {filepath}")


# ─── поиск файла куков ──────────────────────────────────────────────────────
def find_cookies_file() -> str | None:
    user = os.path.expanduser("~")
    search_dirs = [
        os.path.join(user, "Desktop"),
        os.path.join(user, "Desktop", "ytobcoc"),
        os.path.join(user, "Downloads"),
        os.path.join(user, "Downloads", "ytobcoc"),
        os.getcwd(),
    ]
    patterns = ["*youtube*cookie*", "*cookie*youtube*", "cookies*.txt", "*.txt"]
    candidates = []
    for d in search_dirs:
        for pat in patterns:
            for f in glob.glob(os.path.join(d, pat), recursive=False):
                if os.path.isfile(f) and f not in candidates:
                    candidates.append(f)
    if not candidates:
        return None
    candidates.sort(key=os.path.getmtime, reverse=True)
    return candidates[0]


# ─── вспомогательные функции ──────────────────────────────────────────────────
def parse_date(raw) -> str:
    if not raw:
        return ""
    try:
        return datetime.strptime(str(raw), "%Y%m%d").strftime("%Y-%m-%d")
    except ValueError:
        return str(raw)


def wait_for_fix(prompt: str):
    """Выводит сообщение об ошибке и ждёт Enter."""
    print(f"\n{prompt}")
    input("  Исправил? Нажми Enter чтобы продолжить... ")
    print()


def ytdlp_run(args: list[str]) -> tuple[str, str, int]:
    """Запускает yt-dlp и возвращает (stdout, stderr, returncode)."""
    cmd = [sys.executable, "-m", "yt_dlp"] + args
    result = subprocess.run(cmd, capture_output=True, text=True,
                            encoding="utf-8", errors="replace")
    return result.stdout.strip(), result.stderr.strip(), result.returncode


def parse_json_output(stdout: str) -> dict | None:
    if not stdout:
        return None
    try:
        return json.loads(stdout)
    except json.JSONDecodeError:
        for line in reversed(stdout.splitlines()):
            line = line.strip()
            if line.startswith("{"):
                try:
                    return json.loads(line)
                except json.JSONDecodeError:
                    pass
    return None


def is_bot_check(stderr: str) -> bool:
    return "Sign in to confirm" in stderr or "bot" in stderr.lower()


def is_cookies_expired(stderr: str) -> bool:
    return (
        "Sign in to confirm" in stderr
        or "cookies" in stderr.lower() and "authentication" in stderr.lower()
    )


# ─── сохранение в Excel ───────────────────────────────────────────────────────
def save_xlsx(rows: list[dict], output: str):
    """Сохраняет данные в xlsx. Возбуждает PermissionError если файл открыт."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Видео"

    header_labels = ["Заголовок", "Дата", "Просмотры", "Ссылка", "Описание"]
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(bold=True, color="FFFFFF", size=11)

    for col, label in enumerate(header_labels, 1):
        cell = ws.cell(row=1, column=col, value=label)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row_idx, row in enumerate(rows, 2):
        ws.cell(row=row_idx, column=1, value=row["title"])
        ws.cell(row=row_idx, column=2, value=row["published"])
        views = row["views"]
        ws.cell(row=row_idx, column=3, value=int(views) if str(views).isdigit() else views)
        ws.cell(row=row_idx, column=4, value=row["url"])
        ws.cell(row=row_idx, column=5, value=row["description"])

        if row_idx % 2 == 0:
            fill = PatternFill("solid", fgColor="DCE6F1")
            for col in range(1, 6):
                ws.cell(row=row_idx, column=col).fill = fill

    col_widths = [60, 12, 14, 45, 80]
    for col, width in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = width

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=False, vertical="top")
        row[0].alignment = Alignment(wrap_text=True, vertical="top")
        row[4].alignment = Alignment(wrap_text=True, vertical="top")

    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 25

    wb.save(output)  # поднимет PermissionError если файл занят


def save_xlsx_safe(rows: list[dict], output: str):
    """Сохраняет xlsx с повтором если файл открыт в Excel."""
    while True:
        try:
            save_xlsx(rows, output)
            return
        except PermissionError:
            wait_for_fix(
                f"[!] Файл {output} открыт в Excel — закрой его.\n"
                f"    Данные уже собраны и будут сохранены сразу после закрытия."
            )


# ─── получение метаданных с повтором при ошибках ────────────────────────────
def fetch_playlist(url: str, cookies_file: str, max_videos: int | None) -> list[dict]:
    """Получает список видео. Повторяет при ошибке куков."""
    while True:
        flat_args = [
            "--flat-playlist", "--dump-single-json",
            "--quiet", "--no-warnings", "--ignore-errors",
            "--cookies", cookies_file,
        ]
        if max_videos:
            flat_args += ["--playlist-end", str(max_videos)]
        flat_args.append(url)

        stdout, stderr, code = ytdlp_run(flat_args)
        playlist = parse_json_output(stdout)

        if playlist:
            entries = playlist.get("entries") or []
            if entries and isinstance(entries[0], dict) and "entries" in entries[0]:
                entries = entries[0].get("entries") or []
            return [e for e in entries if e and e.get("id")]

        # Диагностика ошибки
        if is_bot_check(stderr) or is_cookies_expired(stderr):
            wait_for_fix(
                "[!] YouTube не принимает куки — сессия устарела.\n"
                "\n"
                "    Что сделать:\n"
                "    1. Открой Chrome и зайди на youtube.com\n"
                "    2. Нажми расширение «Get cookies.txt LOCALLY» → Export\n"
                "    3. Сохрани файл на Рабочий стол\n"
                f"    4. Замени старый файл: {cookies_file}"
            )
            # Ищем обновлённый файл
            fresh = find_cookies_file()
            if fresh and fresh != cookies_file:
                print(f"  Найден новый файл куков: {fresh}")
                cookies_file = fresh
        elif "not found" in stderr.lower() or "no such" in stderr.lower():
            wait_for_fix(
                f"[!] Файл куков не найден: {cookies_file}\n"
                "\n"
                "    Что сделать:\n"
                "    1. Открой Chrome → youtube.com\n"
                "    2. Расширение «Get cookies.txt LOCALLY» → Export\n"
                "    3. Сохрани на Рабочий стол — скрипт найдёт сам"
            )
            found = find_cookies_file()
            if found:
                print(f"  Найден файл куков: {found}")
                cookies_file = found
        elif not stdout:
            wait_for_fix(
                "[!] Не удалось получить список видео.\n"
                "\n"
                f"    Подробности: {stderr[:300] if stderr else 'нет данных'}\n"
                "\n"
                "    Что сделать: проверь ссылку на канал и подключение к интернету."
            )
        else:
            wait_for_fix(
                f"[!] Неожиданная ошибка при загрузке списка видео.\n"
                f"    {stderr[:300]}"
            )


def fetch_video_detail(video_url: str, cookies_file: str) -> dict | None:
    """Получает метаданные одного видео. Возвращает None если видео недоступно."""
    detail_args = [
        "--dump-single-json", "--no-playlist",
        "--quiet", "--no-warnings", "--ignore-errors", "--skip-download",
        "--js-runtimes", "node",
        "--remote-components", "ejs:github",
        "--cookies", cookies_file,
        video_url,
    ]
    stdout, stderr, code = ytdlp_run(detail_args)
    return parse_json_output(stdout)


# ─── основной сбор данных ────────────────────────────────────────────────────
def scrape(channel_url: str, cookies_file: str, max_videos: int | None = None, output_file: str | None = None):
    url = channel_url.rstrip("/")
    if not any(url.endswith(s) for s in ("/videos", "/shorts", "/streams")):
        url += "/videos"

    print(f"\nКанал: {url}")
    print("Загружаю список видео...\n")

    entries = fetch_playlist(url, cookies_file, max_videos)
    if not entries:
        print("[!] Видео не найдены. Проверь ссылку и доступ к каналу.")
        return
    
    total = len(entries)
    print(f"Найдено видео: {total}. Собираю метаданные...\n")

    rows = []
    skipped = 0
    cookies_refreshed = False

    for i, entry in enumerate(entries, 1):
        video_id = entry.get("id")
        if not video_id:
            print(f"[{i}/{total}] [!] ID видео не найден, пропускаю")
            skipped += 1
            continue
            
        video_url = f"https://www.youtube.com/watch?v={video_id}"
        title_preview = (entry.get("title") or video_url)[:72]
        print(f"[{i}/{total}] {title_preview}")

        # Пробуем получить детали — если куки истекли в середине, предлагаем обновить
        for attempt in range(3):
            detail = fetch_video_detail(video_url, cookies_file)
            if detail:
                break

            if attempt < 2:
                # Проверяем напрямую не истекли ли куки
                _, stderr, _ = ytdlp_run([
                    "--dump-single-json", "--no-playlist", "--quiet",
                    "--cookies", cookies_file, video_url
                ])
                if is_bot_check(stderr) and not cookies_refreshed:
                    wait_for_fix(
                        "[!] YouTube снова требует авторизацию — куки устарели в процессе сбора.\n"
                        "\n"
                        "    Что сделать:\n"
                        "    1. Chrome → youtube.com → расширение «Get cookies.txt LOCALLY» → Export\n"
                        "    2. Сохрани на Рабочий стол (заменив старый файл)\n"
                        f"    Текущий файл: {cookies_file}"
                    )
                    fresh = find_cookies_file()
                    if fresh:
                        cookies_file = fresh
                        cookies_refreshed = True
                    break
                else:
                    break
        else:
            detail = None

        if not detail:
            print("  [!] пропускаю")
            skipped += 1
            continue

        desc = (detail.get("description") or "")[:300].replace("\n", " ")
        views = detail.get("view_count")
        rows.append({
            "title": detail.get("title") or entry.get("title") or "",
            "published": parse_date(detail.get("upload_date")),
            "views": str(views) if views is not None else "",
            "url": video_url,
            "description": desc,
        })

    if not rows:
        print("[!] Не удалось собрать ни одного видео.")
        return

    # Определяем имя выходного файла
    if not output_file:
        handle = re.sub(r"[^\w\-]", "_", channel_url.rstrip("/").split("/")[-1])
        output_file = f"{handle}.xlsx"

    print(f"\nСохраняю {output_file}...")
    save_xlsx_safe(rows, output_file)

    abs_output = os.path.abspath(output_file)
    print(f"\n{'='*50}")
    print(f"✓ Собрано видео:  {len(rows)}")
    if skipped:
        print(f"⚠ Пропущено:      {skipped}")
    print(f"📁 Файл:          {abs_output}")
    print(f"{'='*50}")

    # Открываем папку с файлом
    open_file_in_explorer(abs_output)


# ─── запрос куков с обработкой ──────────────────────────────────────────────
def ask_cookies() -> str:
    while True:
        found = find_cookies_file()
        if found:
            print(f"\nНайден файл куков: {found}")
            confirm = input("Использовать его? (Enter = да / n = указать другой): ").strip().lower()
            if confirm != "n":
                return found

        path = input("Путь к cookies.txt: ").strip().strip('"')
        if not path:
            wait_for_fix(
                "[!] Файл куков не указан.\n"
                "\n"
                "    Как получить cookies.txt:\n"
                "    1. Открой Chrome и зайди на youtube.com (войди в аккаунт)\n"
                "    2. Установи расширение «Get cookies.txt LOCALLY»\n"
                "       chrome.google.com/webstore → поиск: get cookies.txt locally\n"
                "    3. Нажми иконку расширения на youtube.com → Export\n"
                "    4. Сохрани файл на Рабочий стол — скрипт найдёт его сам"
            )
            continue

        if os.path.isfile(path):
            return path

        wait_for_fix(
            f"[!] Файл не найден: {path}\n"
            "    Проверь путь или сохрани cookies.txt на Рабочий стол."
        )


# ─── CLI парсер ─────────────────────────────────────────────────────────────
def parse_args():
    parser = argparse.ArgumentParser(
        description="YouTube Channel Scraper — сохраняет видео канала в Excel",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="""
Примеры:
  py youtube_scraper.py
  py youtube_scraper.py --url https://www.youtube.com/@mkbhd
  py youtube_scraper.py --url https://www.youtube.com/@mkbhd --limit 50
  py youtube_scraper.py --url https://www.youtube.com/@mkbhd --limit 50 --output my_videos.xlsx
        """
    )
    parser.add_argument("--url", type=str, help="URL канала YouTube")
    parser.add_argument("--limit", type=int, help="Максимум видео для загрузки")
    parser.add_argument("--output", type=str, help="Путь к выходному файлу Excel")
    parser.add_argument("--cookies", type=str, help="Путь к файлу cookies.txt")
    
    return parser.parse_args()


# ─── точка входа ────────────────────────────────────────────────────────────
def main():
    print("=" * 50)
    print("  YouTube Channel Scraper → Excel")
    print("=" * 50)

    args = parse_args()

    # Получаем URL канала
    if args.url:
        channel_url = args.url
    else:
        channel_url = input("\nСсылка на канал: ").strip().lstrip("﻿")
        if not channel_url:
            print("Ссылка не указана.")
            sys.exit(1)

    # Получаем лимит видео
    if args.limit:
        max_videos = args.limit
    else:
        limit_str = input("Сколько последних видео собрать? (Enter = все): ").strip()
        max_videos = int(limit_str) if limit_str.isdigit() else None

    # Получаем путь к куки файлу
    if args.cookies:
        cookies_file = args.cookies
        if not os.path.isfile(cookies_file):
            print(f"[!] Файл куков не найден: {cookies_file}")
            sys.exit(1)
    else:
        cookies_file = ask_cookies()

    # Запускаем скрейпинг
    scrape(channel_url, cookies_file, max_videos, args.output)


if __name__ == "__main__":
    main()
